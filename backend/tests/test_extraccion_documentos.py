"""extraer_texto contra archivos .txt/.md/.pdf/.docx reales, incluido el caso
"no hay texto legible" (por ejemplo, un PDF sin capa de texto).
"""

from __future__ import annotations

import io

import pytest
from docx import Document
from pypdf import PdfWriter

from app.adaptadores.extraccion_documentos import extraer_texto
from app.dominio.errores import DocumentoNoLegible

TEXTO_LARGO = "Especificación de prueba con contenido suficiente para pasar el mínimo. " * 2


def test_extrae_texto_plano() -> None:
    assert extraer_texto("spec.txt", TEXTO_LARGO.encode("utf-8")) == TEXTO_LARGO.strip()


def test_extrae_markdown() -> None:
    contenido = f"# Título\n\n{TEXTO_LARGO}".encode()
    assert "# Título" in extraer_texto("spec.md", contenido)


def test_formato_no_soportado_da_documento_no_legible() -> None:
    with pytest.raises(DocumentoNoLegible):
        extraer_texto("spec.png", b"no importa el contenido")


def test_texto_demasiado_corto_da_documento_no_legible() -> None:
    with pytest.raises(DocumentoNoLegible):
        extraer_texto("spec.txt", b"hola")


def test_extrae_docx_con_parrafos_reales() -> None:
    documento = Document()
    documento.add_paragraph("Especificación funcional")
    documento.add_paragraph(TEXTO_LARGO)
    buffer = io.BytesIO()
    documento.save(buffer)

    texto = extraer_texto("spec.docx", buffer.getvalue())

    assert "Especificación funcional" in texto


def test_docx_vacio_da_documento_no_legible() -> None:
    documento = Document()
    buffer = io.BytesIO()
    documento.save(buffer)

    with pytest.raises(DocumentoNoLegible):
        extraer_texto("spec.docx", buffer.getvalue())


def test_pdf_sin_capa_de_texto_da_documento_no_legible() -> None:
    """Una página en blanco simula un PDF escaneado sin texto extraíble."""
    escritor = PdfWriter()
    escritor.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    escritor.write(buffer)

    with pytest.raises(DocumentoNoLegible):
        extraer_texto("spec.pdf", buffer.getvalue())
